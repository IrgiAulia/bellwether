#!/usr/bin/env python3
"""Generate a data-extraction form (CSV/XLSX) + data dictionary from modular column sets.

Usage:
  python extraction_template.py --modules core,quantitative_binary,agri --out extraction \
      [--extra "soil_type;biochar_feedstock"] [--xlsx]
  python extraction_template.py --list

Modules:  core | quantitative_binary | quantitative_continuous | quantitative_correlation | qualitative |
          intervention_tidier | agri | policy_social | education | environment | clinical | mixed_methods
Writes <out>.csv (one row per study-result), <out>_dictionary.csv (variable, definition, allowed values) and,
with --xlsx, <out>.xlsx (form + dictionary sheets with dropdowns for coded columns).
PRISMA items: 9 (how data collected), 10a/10b (data items, assumptions), 17, 19, 27 (share form + dictionary).
Rules baked in: 'NR' = not reported (never leave blank or guess); record WHERE each number came from
(page/table/figure); record who extracted and who verified; keep one row per independent comparison.
"""
import argparse, csv, sys

M = {}
M["core"] = [
    ("study_id", "Unique study ID (first author + year; add a/b if needed). One STUDY may have several reports.", ""),
    ("report_ids", "All reports (record_ids) that belong to this study; mark the primary report", ""),
    ("citation", "Full citation of primary report", ""),
    ("doi", "DOI of primary report", ""),
    ("country_region", "Country/region where the study was done", ""),
    ("setting_context", "Setting/context (e.g. agro-ecological zone, health system level, regime type, school type)", ""),
    ("study_design", "Design as reported by authors AND as classified by reviewers", "RCT|cluster RCT|non-randomised trial|cohort|case-control|cross-sectional|quasi-experiment|field experiment|pot experiment|case study|qualitative|mixed methods|other"),
    ("study_period", "Data collection period / duration / follow-up", ""),
    ("sample_size_total", "Total N analysed (units: participants, plots, farms, countries, documents...)", ""),
    ("unit_of_analysis", "Unit analysed (individual, plot, farm, district, country, policy episode...)", ""),
    ("population_description", "Population/subject characteristics relevant to eligibility & moderators (age, crop/cultivar, soil, actors...)", ""),
    ("phenomenon_intervention_exposure", "Intervention / exposure / phenomenon of interest as delivered", ""),
    ("comparator_context", "Comparator / control / counterfactual", ""),
    ("outcomes_measured", "Outcomes measured (domain, instrument/metric, unit, direction of benefit)", ""),
    ("timepoints", "Time point(s) of outcome measurement", ""),
    ("funding_source", "Funding source(s)", ""),
    ("competing_interests", "Author competing interests declared", ""),
    ("data_source_for_extraction", "Where numbers came from: table/figure/text/supplement/author contact (with page)", ""),
    ("assumptions_or_imputations", "Any assumption made about missing/unclear information (PRISMA 10b) e.g. imputed SD, assumed age range", ""),
    ("notes", "Anything else", ""),
    ("extracted_by", "Initials of extractor", ""),
    ("verified_by", "Initials of second extractor/verifier", ""),
    ("discrepancy_resolution", "How disagreements were resolved (discussion/third reviewer)", ""),
]
M["quantitative_binary"] = [
    ("comparison_id", "ID of the specific comparison/outcome (one row per independent comparison)", ""),
    ("outcome_label", "Outcome label as defined in the review", ""),
    ("events_group1", "Participants/units with the event, group 1 (intervention/exposed)", "integer|NR"),
    ("total_group1", "Total analysed in group 1", "integer|NR"),
    ("events_group2", "Participants/units with the event, group 2 (control)", "integer|NR"),
    ("total_group2", "Total analysed in group 2", "integer|NR"),
    ("effect_estimate_reported", "Effect as reported by authors (RR/OR/HR + scale)", ""),
    ("ci_lower", "Lower 95% CI", ""), ("ci_upper", "Upper 95% CI", ""),
    ("adjusted", "Adjusted for covariates? which", "yes|no|NR"),
    ("clustering_adjusted", "Cluster-adjusted? design effect/ICC used", "yes|no|not applicable|NR"),
]
M["quantitative_continuous"] = [
    ("comparison_id", "ID of the specific comparison/outcome", ""),
    ("outcome_label", "Outcome label as defined in the review", ""),
    ("instrument_unit", "Instrument/scale and unit; range; direction of benefit (higher=better?)", ""),
    ("mean_group1", "Mean (or change from baseline) group 1", "number|NR"),
    ("sd_group1", "SD group 1 (convert SE/CI to SD and note in assumptions)", "number|NR"),
    ("n_group1", "N analysed group 1", "integer|NR"),
    ("mean_group2", "Mean group 2", "number|NR"), ("sd_group2", "SD group 2", "number|NR"),
    ("n_group2", "N analysed group 2", "integer|NR"),
    ("sd_type_reported", "What the paper labelled SD/SE/CI/IQR/range (never assume)", "SD|SE|95% CI|IQR|range|NR"),
    ("median_iqr_used", "Median/IQR reported instead of mean/SD? (conversion needs a documented method)", "yes|no"),
    ("value_type", "Final values or change scores", "final|change|NR"),
]
M["quantitative_correlation"] = [
    ("comparison_id", "ID of the relationship", ""), ("relationship", "X -> Y relationship (variables, direction)", ""),
    ("r", "Correlation coefficient (type: Pearson/Spearman)", "number|NR"), ("n", "N for that correlation", "integer|NR"),
    ("controls", "Partial correlation/regression? controls", ""),
]
M["qualitative"] = [
    ("qual_method", "Data collection method(s)", "interviews|focus groups|ethnography|open-ended survey|documents|observation|other"),
    ("qual_analysis", "Analytic approach claimed by authors", "thematic|grounded theory|phenomenology|content analysis|framework|discourse|other|NR"),
    ("qual_sampling", "Sampling strategy & sample size (participants)", ""),
    ("theory_framework", "Theoretical/philosophical framework used by authors", ""),
    ("reflexivity", "Researcher position/reflexivity reported", "yes|partial|no"),
    ("key_themes_authors", "Main themes/concepts (authors' own words, with page)", ""),
    ("illustrative_quotes", "Participant quotations supporting each theme (page)", ""),
    ("data_richness", "Thick/thin data (for CERQual adequacy)", "thick|moderate|thin"),
]
M["intervention_tidier"] = [
    ("tidier_what", "What: materials & procedures", ""), ("tidier_who", "Who delivered", ""),
    ("tidier_how", "How: mode of delivery", ""), ("tidier_where", "Where", ""),
    ("tidier_when_howmuch", "When and how much (dose, frequency, duration)", ""),
    ("tidier_tailoring", "Tailoring/modifications", ""), ("tidier_fidelity", "Fidelity/adherence", ""),
]
M["agri"] = [
    ("crop_or_livestock", "Species/cultivar/breed", ""), ("agro_ecological_zone", "Climate/AEZ (e.g. Köppen), altitude, rainfall", ""),
    ("soil_properties", "Soil type & baseline properties (pH, SOC, texture)", ""),
    ("management_practice", "Practice details: rate, timing, method, feedstock/variety/dose", ""),
    ("experimental_unit", "Plot size, replicates, design (RCBD/split-plot), pot vs field", ""),
    ("season_year", "Season(s)/years; number of site-years", ""),
    ("yield_or_response_metric", "Response variable & unit (t/ha, kg/plant, %)", ""),
    ("control_type", "Control definition (no input, farmer practice, standard fertiliser)", ""),
    ("shared_control_flag", "Does this row share a control with another row? (non-independence)", "yes|no"),
]
M["policy_social"] = [
    ("actors_units", "Actors/institutions/units studied", ""), ("policy_institution", "Policy / institutional feature / reform", ""),
    ("regime_context", "Political/institutional context (regime type, decentralisation, period)", ""),
    ("identification_strategy", "Causal identification strategy", "RCT|natural experiment|DiD|RD|IV|synthetic control|panel FE|cross-section|process tracing|comparative case|descriptive|other"),
    ("outcome_metric", "Outcome operationalisation (e.g. turnout %, index score)", ""),
    ("effect_scale", "Reported effect scale (coef, SE, partial r, marginal effect)", ""),
    ("mechanism_reported", "Mechanism/theory of change proposed", ""),
]
M["education"] = [
    ("learner_level", "Learner level/age/grade", ""), ("delivery_mode", "Delivery mode (face-to-face/online/blended)", ""),
    ("instrument_validity", "Outcome instrument & validity/reliability evidence", ""), ("dosage", "Sessions x duration", ""),
]
M["environment"] = [
    ("ecosystem_habitat", "Ecosystem/habitat/biome", ""), ("spatial_scale", "Spatial scale & extent", ""),
    ("pressure_exposure", "Pressure/exposure & intensity", ""), ("response_variable", "Response variable & unit", ""),
    ("baseline_reference", "Reference/baseline condition", ""),
]
M["clinical"] = [
    ("condition", "Condition/disease & severity", ""), ("age_mean_sd", "Age mean (SD)", ""), ("percent_female", "% female", ""),
    ("dose_regimen", "Dose/regimen", ""), ("adverse_events_definition", "Harms definition & how ascertained", ""),
    ("registration_id", "Trial registration ID (for selective-reporting checks)", ""),
]
M["mixed_methods"] = [
    ("qual_component", "Qualitative component summary", ""), ("quant_component", "Quantitative component summary", ""),
    ("integration_point", "How/where components were integrated by authors", ""),
]


def build(mods, extra):
    cols = []
    for m in mods:
        if m not in M:
            sys.exit(f"Unknown module {m}. Options: {', '.join(M)}")
        cols += [c for c in M[m] if c[0] not in [x[0] for x in cols]]
    for e in extra:
        cols.append((e, "Custom variable (define!)", ""))
    return cols


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--modules", default="core"); ap.add_argument("--out")
    ap.add_argument("--extra", default=""); ap.add_argument("--xlsx", action="store_true"); ap.add_argument("--list", action="store_true")
    a = ap.parse_args()
    if a.list:
        for k, v in M.items():
            print(f"{k}: {len(v)} vars")
        return
    if not a.out:
        ap.error("--out required")
    mods = [m.strip() for m in a.modules.split(",") if m.strip()]
    if "core" not in mods:
        mods.insert(0, "core")
    cols = build(mods, [e.strip() for e in a.extra.split(";") if e.strip()])
    with open(a.out + ".csv", "w", newline="", encoding="utf-8-sig") as fh:
        csv.writer(fh).writerow([c[0] for c in cols])
    with open(a.out + "_dictionary.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["variable", "definition", "allowed_values"])
        w.writerows(cols)
    if a.xlsx:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "extraction"
        ws.append([c[0] for c in cols])
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDEBF7")
            ws.column_dimensions[cell.column_letter].width = max(14, min(40, len(c[0]) + 4))
            opts = [o for o in c[2].split("|") if o and o not in ("integer", "number")]
            if opts and len(",".join(opts)) < 250:
                dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
                ws.add_data_validation(dv); dv.add(f"{cell.column_letter}2:{cell.column_letter}500")
        ws.freeze_panes = "B2"
        wd = wb.create_sheet("dictionary"); wd.append(["variable", "definition", "allowed_values"])
        for c in cols:
            wd.append(list(c))
        wd.column_dimensions["A"].width = 32; wd.column_dimensions["B"].width = 100; wd.column_dimensions["C"].width = 60
        wb.save(a.out + ".xlsx")
    print(f"{len(cols)} variables -> {a.out}.csv, {a.out}_dictionary.csv" + (f", {a.out}.xlsx" if a.xlsx else ""))


if __name__ == "__main__":
    main()
