#!/usr/bin/env python3
"""Heuristic completeness check of a review report against PRISMA 2020 items (27 items / 42 sub-items).

Usage:
  python prisma_lint.py --report report.md [--out prisma_checklist.csv]
  (accepts .md, .txt, .docx)

It searches for keyword evidence per item and prints a checklist with the first matching line. It CANNOT judge
quality: 'FOUND' means a relevant term exists, 'MISSING' means nothing relevant was detected -> read the report
and fix or mark 'not applicable' with a reason. Always finish by a manual pass with references/reporting-prisma2020.md.
"""
import argparse, csv, re, sys

ITEMS = [
 ("1", "Title", "Identify the report as a systematic review", [r"systematic (literature )?review", r"tinjauan (pustaka|literatur) sistematis", r"meta-?(analysis|synthesis|sintesis)"], "title"),
 ("2", "Abstract", "Structured abstract (PRISMA 2020 for Abstracts, 12 items)", [r"abstract|abstrak", r"registration|PROSPERO|OSF|registered"], None),
 ("3", "Introduction", "Rationale in context of existing knowledge", [r"rationale|background|latar belakang|existing (review|evidence)|knowledge gap|research gap"], None),
 ("4", "Introduction", "Explicit objectives / questions (PICO or other framework)", [r"objective|aim of|research question|tujuan|pertanyaan penelitian", r"PICO|PECO|SPIDER|PCC|PICo|SPICE|CIMO|ECLIPSE|framework"], None),
 ("5", "Methods", "Eligibility criteria & grouping for synthesis", [r"eligibility|inclusion criteria|exclusion criteria|kriteria inklusi|kriteria eksklusi"], None),
 ("6", "Methods", "Information sources + date last searched", [r"last searched|date of (the )?search|searched (on|in)|terakhir (dicari|ditelusuri)", r"Scopus|Web of Science|PubMed|MEDLINE|Embase|database"], None),
 ("7", "Methods", "Full search strategies", [r"search (strategy|string)|strategi pencarian|kata kunci|Boolean", r"appendix|supplement|lampiran"], None),
 ("8", "Methods", "Selection process (reviewers, independence, automation)", [r"screen(ed|ing)", r"independent(ly)?|two reviewers|second reviewer|dual|reviewer"], None),
 ("9", "Methods", "Data collection process", [r"data (extraction|collection)|extracted|ekstraksi data", r"form|template|piloted|two reviewers|verified"], None),
 ("10a", "Methods", "Outcomes sought & how results were selected", [r"outcome (domain|measure)s?|primary outcome|luaran|outcomes? (were )?(sought|defined)"], None),
 ("10b", "Methods", "Other variables + assumptions about missing info", [r"data items?|variables? (extracted|collected)|assum(ed|ption)|not reported|NR\b|missing information"], None),
 ("11", "Methods", "Risk of bias / critical appraisal methods", [r"risk of bias|RoB ?2|ROBINS|Newcastle|CASP|JBI|critical appraisal|MMAT|quality assessment|penilaian kualitas"], None),
 ("12", "Methods", "Effect measures", [r"risk ratio|odds ratio|mean difference|SMD|Hedges|response ratio|correlation|effect (measure|size)|hazard ratio|ukuran efek"], None),
 ("13a", "Methods", "How studies were eligible for each synthesis", [r"grouped|grouping|eligible for (each )?synthesis|tabulat|synthesis (groups?|categories)|dikelompokkan"], None),
 ("13b", "Methods", "Data preparation (conversions, missing statistics)", [r"convert(ed|sion)|imput(ed|ation)|continuity correction|standard (deviation|error)|transform"], None),
 ("13c", "Methods", "Tabulation / visual display methods", [r"forest plot|funnel|table(s)? (of|summar)|visual(ly)? (display|present)|figure|matrix"], None),
 ("13d", "Methods", "Synthesis methods, model, heterogeneity, software", [r"random-effects|fixed-effect|meta-analy|thematic synthesis|meta-ethnograph|meta-aggregat|narrative synthesis|SWiM|framework synthesis|vote counting", r"software|R \(|metafor|Stata|RevMan|NVivo|ATLAS|Python|version"], None),
 ("13e", "Methods", "Exploring causes of heterogeneity", [r"subgroup|meta-regression|moderator|heterogeneity"], None),
 ("13f", "Methods", "Sensitivity analyses", [r"sensitivity analys|leave-one-out|robustness"], None),
 ("14", "Methods", "Reporting bias assessment methods", [r"publication bias|reporting bias|funnel|Egger|ROB-ME|selective (outcome )?reporting|missing results"], None),
 ("15", "Methods", "Certainty of evidence methods", [r"GRADE|CERQual|certainty of (the )?evidence|confidence in (the )?(evidence|findings)|tingkat keyakinan"], None),
 ("16a", "Results", "Search & selection results + flow diagram", [r"flow (diagram|chart)|PRISMA (flow|diagram)|records (identified|screened)|diagram alir"], None),
 ("16b", "Results", "Excluded near-miss studies with reasons", [r"excluded (studies|reports)|reasons? for exclusion|alasan eksklusi|exclusion reasons?"], None),
 ("17", "Results", "Study characteristics table + citations", [r"characteristics of (the )?(included )?stud|table\s*\d.*characteristics|karakteristik studi"], None),
 ("18", "Results", "Risk of bias in studies (per domain)", [r"risk of bias (in|of|for|assessment|results)|appraisal (results|of included)|traffic|domain"], None),
 ("19", "Results", "Results of individual studies", [r"individual stud(y|ies) results?|study-level|forest plot|effect estimate|per study"], None),
 ("20a", "Results", "Characteristics/risk of bias among studies in each synthesis", [r"contributing (to|studies)|studies (included|contributing) in (the )?(each )?(synthesis|meta-analysis)"], None),
 ("20b", "Results", "Results of syntheses (estimate, CI, heterogeneity, direction)", [r"95% ?CI|confidence interval|pooled|I\s?2|I²|tau|prediction interval|themes?|synthesi[sz]ed finding", r"favou?r"], None),
 ("20c", "Results", "Results of heterogeneity investigations", [r"subgroup|meta-regression|moderator|heterogeneity|context(ual)? (factor|variation)"], None),
 ("20d", "Results", "Results of sensitivity analyses", [r"sensitivity analys|leave-one-out|robust"], None),
 ("21", "Results", "Risk of bias due to missing results", [r"funnel|Egger|publication bias|reporting bias|small-study"], None),
 ("22", "Results", "Certainty of evidence per outcome", [r"GRADE|CERQual|certainty|Summary of Findings|summary of findings|SoF"], None),
 ("23a", "Discussion", "Interpretation in context of other evidence", [r"discussion|consistent with|in line with|compared with (previous|other)|dibandingkan"], None),
 ("23b", "Discussion", "Limitations of the evidence", [r"limitation(s)? of the (included )?evidence|limitations? (of|in) the (included )?studies|evidence (was|is) limited|keterbatasan bukti"], None),
 ("23c", "Discussion", "Limitations of review processes", [r"limitation(s)? of (this|the) review|review process limitation|strengths and limitations|keterbatasan (penelitian|review)"], None),
 ("23d", "Discussion", "Implications for practice, policy, research", [r"implication(s)? for (practice|policy|research)|future research|implikasi"], None),
 ("24a", "Other", "Registration information", [r"PROSPERO|CRD\d{6,}|OSF|registered|registration|not registered|tidak terdaftar|registrasi"], None),
 ("24b", "Other", "Protocol accessibility", [r"protocol"], None),
 ("24c", "Other", "Amendments / deviations from protocol", [r"amendment|deviation|differences? from (the )?protocol|changes? (to|from) (the )?protocol|perubahan protokol"], None),
 ("25", "Other", "Support / funding & role of funders", [r"funding|funder|financial support|grant|pendanaan"], None),
 ("26", "Other", "Competing interests", [r"competing interest|conflict of interest|conflicts? of interest|konflik kepentingan"], None),
 ("27", "Other", "Availability of data, code, materials", [r"data availability|availability of data|OSF|Zenodo|figshare|Dryad|repository|code availability|materials (are|is) available|ketersediaan data"], None),
]


def read_report(p):
    if p.lower().endswith(".docx"):
        import docx
        d = docx.Document(p)
        lines = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return lines
    txt = open(p, encoding="utf-8", errors="replace").read()
    # template guidance lives in HTML comments: it must not count as evidence (keep line numbers)
    txt = re.sub(r"<!--.*?-->", lambda m: "\n" * m.group(0).count("\n"), txt, flags=re.S)
    return txt.splitlines()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--report", required=True); ap.add_argument("--out")
    a = ap.parse_args()
    lines = read_report(a.report)
    text_all = "\n".join(lines)
    rows, miss = [], []
    for iid, sec, name, pats, where in ITEMS:
        hits = []
        for pat in pats:
            m = None
            search_lines = lines[:6] if where == "title" else lines
            for i, ln in enumerate(search_lines):
                if re.search(pat, ln, re.I):
                    m = (i + 1, ln.strip()[:110]); break
            hits.append(m)
        status = "FOUND" if all(hits) else ("PARTIAL" if any(hits) else "MISSING")
        first = next((h for h in hits if h), None)
        rows.append({"item": iid, "section": sec, "requirement": name, "status": status,
                     "evidence_line": f"L{first[0]}: {first[1]}" if first else "", "manual_check/location": "", "not_applicable_reason": ""})
        if status != "FOUND":
            miss.append((iid, name, status))
    if a.out:
        with open(a.out, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    n_found = sum(1 for r in rows if r["status"] == "FOUND")
    print(f"PRISMA 2020 heuristic check: {n_found}/{len(rows)} sub-items with evidence; {len(rows) - n_found} to review\n")
    for iid, name, st in miss:
        print(f"  [{st:<7}] {iid:<4} {name}")
    words = len(re.findall(r"\w+", text_all))
    print(f"\n(report length ~{words} words). Reminder: FOUND != adequate. Check essential elements manually.")


if __name__ == "__main__":
    main()
